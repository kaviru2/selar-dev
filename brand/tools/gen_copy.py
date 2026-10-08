"""Write the taglines / pitch / captions doc (Markdown) and a DOCX copy from copy_text.py."""
import os, sys
sys.path.insert(0, os.path.dirname(__file__))
import copy_text as C

OUT = os.path.join(os.path.dirname(__file__), "..", "out")

ALT = [
    ("OG / Twitter card", "SELAR research prototype: Linny the linnet links a passage you are reading to one you read earlier. AI suggests, sources show, you explain and decide."),
    ("Carousel 1", "Linny the linnet thinking, next to the question: ever read something and thought, wait, I've seen this before?"),
    ("Carousel 2", "Four steps: read, notice a suggested link, compare both passages, explain and decide."),
    ("Carousel 3", "Survey results: 31 of 36 wanted a linking tool; 28 of 36 worried about AI accuracy; 24 of 36 worried about less active thinking."),
    ("Square 1", "The SELAR principle on green: AI suggests. Sources show. You decide."),
    ("Square 2", "Linny the linnet holding a page, introducing the SELAR mascot."),
    ("Flyer", "A4 poster inviting students to try the SELAR research prototype and share feedback, with a QR code to selar-console.vercel.app."),
]
NAMES = {"carousel": "Carousel (3 slides, 1080x1350): carousel-01..03.png",
         "square-principle": "Square post: square-01-principle.png",
         "square-linny": "Square post: square-02-linny.png"}


def sections():
    yield ("h1", "SELAR: taglines, pitch and social captions")
    yield ("p", "Research prototype · UCSC IS4101 Group 04 · v1.0, October 2026. Drafts only: nothing here has been posted.")
    yield ("h2", "Taglines (recommended first)")
    for i, (t, n) in enumerate(C.TAGLINES, 1):
        yield ("item", (f"{i}. {t}", n))
    yield ("h2", f"30-second elevator pitch ({len(C.PITCH.split())} words)")
    yield ("p", C.PITCH)
    yield ("h2", f"Longer pitch, about 60 seconds ({len(C.PITCH_LONG.split())} words)")
    yield ("p", C.PITCH_LONG)
    yield ("h2", "Draft captions (team review before posting)")
    for k, v in C.CAPTIONS.items():
        yield ("h3", NAMES[k])
        for para in v.split("\n\n"):
            yield ("p", para)
    yield ("h2", "Alt text")
    for k, v in ALT:
        yield ("bullet", f"{k}: {v}")
    yield ("h2", "Claim rules (apply to every edit)")
    for r in C.CLAIM_RULES:
        yield ("bullet", r)
    yield ("h2", "Survey wording")
    yield ("p", C.SURVEY_NOTE + " " + " ".join(f"{n}: {d}." for n, d in C.SURVEY))


def md():
    out = []
    for kind, v in sections():
        if kind == "h1": out += [f"# {v}", ""]
        elif kind == "h2": out += [f"## {v}", ""]
        elif kind == "h3": out += [f"### {v}", ""]
        elif kind == "p": out += [v, ""]
        elif kind == "bullet": out += [f"- {v}"]
        elif kind == "item": out += [f"**{v[0]}**  ", f"{v[1]}", ""]
    open(os.path.join(OUT, "SELAR-taglines-pitch-captions.md"), "w").write("\n".join(out).replace("\n## ", "\n\n## ") + "\n")


def docx():
    from docx import Document
    from docx.shared import Pt, RGBColor
    d = Document()
    st = d.styles["Normal"]; st.font.name = "Figtree"; st.font.size = Pt(11)
    for kind, v in sections():
        if kind in ("h1", "h2", "h3"):
            h = d.add_heading(v, level={"h1": 0, "h2": 1, "h3": 2}[kind])
            for r in h.runs: r.font.color.rgb = RGBColor(0x1F, 0x51, 0x35)
        elif kind == "p": d.add_paragraph(v)
        elif kind == "bullet": d.add_paragraph(v, style="List Bullet")
        elif kind == "item":
            p = d.add_paragraph(); r = p.add_run(v[0]); r.bold = True
            p.add_run("\n" + v[1])
    d.save(os.path.join(OUT, "SELAR-taglines-pitch-captions.docx"))


if __name__ == "__main__":
    md()
    try:
        docx(); print("docx ok")
    except ImportError as e:
        print("no python-docx:", e)
